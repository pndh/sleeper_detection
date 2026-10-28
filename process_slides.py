#!/usr/bin/env python3
"""
Update authentic presentation deck from /home/pndhpndh/Downloads/AIA Presentation slides.pptx
with contact emails for Raghav Agarwal and Huy Pham, then export to PDF, PNG slides, and artifacts.
"""
import os
import shutil
import subprocess
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

SRC_PPTX = "/home/pndhpndh/Downloads/AIA Presentation slides.pptx"
OUT_PPTX_NAME = "AIA Presentation slides.pptx"
LEGACY_PPTX_NAME = "AIA_Illinois_Midpoint_Checkpoint_Raghav_Agarwal.pptx"

OUT_PDF_NAME = "AIA Presentation slides.pdf"
LEGACY_PDF_NAME = "AIA_Illinois_Midpoint_Checkpoint_Raghav_Agarwal.pdf"

ARTIFACT_DIR = "/home/pndhpndh/.gemini/antigravity/brain/fd4b0ee4-31d9-430c-af29-2bcb38cd9d97"
SITE_DOWNLOADS = "/home/pndhpndh/sleeper_detection/site/public/downloads"
SITE_SLIDES = "/home/pndhpndh/sleeper_detection/site/public/slides"
RESULTS_SLIDES = "/home/pndhpndh/sleeper_detection/results/presentation_slides"

def update_presentation():
    print(f"Loading {SRC_PPTX}...")
    prs = Presentation(SRC_PPTX)
    print(f"Loaded presentation with {len(prs.slides)} slides.")

    # 1. Update Slide 1 (Shape 7):
    # Original:
    #   Raghav Agarwal  ·  Huy Pham
    #   Mentors: Enyi Jiang, Himnish Hunma  ·  October 2026
    s1_shape = prs.slides[0].shapes[7]
    tf1 = s1_shape.text_frame
    tf1.word_wrap = True
    
    # Check paragraphs
    p0 = tf1.paragraphs[0]
    p0.text = "Raghav Agarwal (raghava7@illinois.edu)  ·  Huy Pham (huynp2@illinois.edu)"
    if p0.runs:
        p0.runs[0].font.name = "Arial"
        p0.runs[0].font.size = Pt(17)
        p0.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    if len(tf1.paragraphs) > 1:
        p1 = tf1.paragraphs[1]
        p1.text = "Mentors: Enyi Jiang, Himnish Hunma  ·  October 2026"
        if p1.runs:
            p1.runs[0].font.name = "Arial"
            p1.runs[0].font.size = Pt(17)
            p1.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # 2. Update Slide 14 (Shape 8):
    # Original:
    #   Code, reports, features and labels on GitHub (RaghavAgarwal08/sleeper_detection, branch raghav/probe-live-inert). Controls rerun on a laptop in 2 minutes
    s14_shape = prs.slides[13].shapes[8]
    tf14 = s14_shape.text_frame
    tf14.word_wrap = True
    p14_0 = tf14.paragraphs[0]
    p14_0.text = "Code, reports, features and labels on GitHub (RaghavAgarwal08/sleeper_detection, branch raghav/probe-live-inert). Controls rerun on a laptop in 2 minutes."
    if p14_0.runs:
        p14_0.runs[0].font.name = "Arial"
        p14_0.runs[0].font.size = Pt(12)
        p14_0.runs[0].font.color.rgb = RGBColor(0x13, 0x29, 0x4B)

    # Add contact paragraph
    p14_1 = tf14.add_paragraph()
    p14_1.text = "Contact: Raghav Agarwal (raghava7@illinois.edu)  ·  Huy Pham (huynp2@illinois.edu)"
    p14_1.space_before = Pt(6)
    if p14_1.runs:
        p14_1.runs[0].font.name = "Arial"
        p14_1.runs[0].font.size = Pt(12)
        p14_1.runs[0].font.bold = True
        p14_1.runs[0].font.color.rgb = RGBColor(0xFF, 0x5F, 0x05)  # Illinois Orange highlight for contact

    # Save to workspace root
    out_path = f"/home/pndhpndh/sleeper_detection/{OUT_PPTX_NAME}"
    legacy_path = f"/home/pndhpndh/sleeper_detection/{LEGACY_PPTX_NAME}"
    prs.save(out_path)
    shutil.copyfile(out_path, legacy_path)
    print(f"Saved {out_path} and {legacy_path}")

    # Copy to downloads & artifact dir
    os.makedirs(SITE_DOWNLOADS, exist_ok=True)
    shutil.copyfile(out_path, os.path.join(SITE_DOWNLOADS, OUT_PPTX_NAME))
    shutil.copyfile(out_path, os.path.join(SITE_DOWNLOADS, LEGACY_PPTX_NAME))
    shutil.copyfile(out_path, os.path.join(ARTIFACT_DIR, OUT_PPTX_NAME))
    shutil.copyfile(out_path, os.path.join(ARTIFACT_DIR, LEGACY_PPTX_NAME))
    print("Copied PPTX to downloads and artifacts.")

    # Convert to PDF via libreoffice
    cmd = [
        "libreoffice", "--headless", "--convert-to", "pdf",
        out_path, "--outdir", "/home/pndhpndh/sleeper_detection"
    ]
    print("Converting PPTX to PDF with LibreOffice...")
    subprocess.run(cmd, check=True)
    pdf_path = f"/home/pndhpndh/sleeper_detection/{OUT_PDF_NAME}"
    legacy_pdf_path = f"/home/pndhpndh/sleeper_detection/{LEGACY_PDF_NAME}"
    shutil.copyfile(pdf_path, legacy_pdf_path)

    # Copy PDF to downloads & artifact dir
    shutil.copyfile(pdf_path, os.path.join(SITE_DOWNLOADS, OUT_PDF_NAME))
    shutil.copyfile(pdf_path, os.path.join(SITE_DOWNLOADS, LEGACY_PDF_NAME))
    shutil.copyfile(pdf_path, os.path.join(ARTIFACT_DIR, OUT_PDF_NAME))
    shutil.copyfile(pdf_path, os.path.join(ARTIFACT_DIR, LEGACY_PDF_NAME))
    print("Copied PDF to downloads and artifacts.")

    # Render PNG slides
    os.makedirs(SITE_SLIDES, exist_ok=True)
    os.makedirs(RESULTS_SLIDES, exist_ok=True)
    # Clear old slides first
    for f in os.listdir(SITE_SLIDES):
        if f.startswith("slide-") and f.endswith(".png"):
            os.remove(os.path.join(SITE_SLIDES, f))
    for f in os.listdir(RESULTS_SLIDES):
        if f.startswith("slide-") and f.endswith(".png"):
            os.remove(os.path.join(RESULTS_SLIDES, f))

    prefix = os.path.join(SITE_SLIDES, "slide")
    render_cmd = ["pdftoppm", "-png", "-r", "150", pdf_path, prefix]
    print("Rendering slides to PNG with pdftoppm...")
    subprocess.run(render_cmd, check=True)

    # Copy to RESULTS_SLIDES
    for f in sorted(os.listdir(SITE_SLIDES)):
        if f.startswith("slide-") and f.endswith(".png"):
            shutil.copyfile(os.path.join(SITE_SLIDES, f), os.path.join(RESULTS_SLIDES, f))
    print(f"Generated {len(os.listdir(SITE_SLIDES))} slides in {SITE_SLIDES} and {RESULTS_SLIDES}.")

if __name__ == "__main__":
    update_presentation()
